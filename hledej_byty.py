"""
Hledač bytů: Sreality + Bezrealitky + Bazoš -> web/byty.json (+ byty.xlsx)

Spuštění:  python hledej_byty.py
Každé spuštění stáhne aktuální nabídky, porovná je s databází (byty.db),
zapíše data pro webový přehled (web/index.html čte web/byty.json)
a lokálně i Excel s listy "Nové" (poprvé viděné při tomto běhu) a "Všechny" (aktivní).
Na GitHubu (proměnná CI) se Excel nevytváří.
"""
import json
import math
import os
import re
import sqlite3
import sys
import time
import unicodedata
from datetime import datetime
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

# ---------------------------------------------------------------- NASTAVENÍ --
MAX_CENA = 22_000            # nájem VČETNĚ poplatků, Kč (nájem samotný taky nesmí být vyšší)
POCET_OSOB = 3               # poplatky uvedené "za osobu" se násobí tímto číslem
MAX_VZDALENOST_KM = 12       # vzdušnou čarou od práce; None = celá Praha
PRACE = (50.1185, 14.4100)   # Zoo Praha / Botanická zahrada, Troja
PAUZA_S = 0.7                # pauza mezi požadavky na detail (šetrnost k webům)

SLOZKA = Path(__file__).resolve().parent
DB_SOUBOR = SLOZKA / "byty.db"
EXCEL_SOUBOR = SLOZKA / "byty.xlsx"
LOG_SOUBOR = SLOZKA / "log.txt"
WEB_JSON = SLOZKA / "web" / "byty.json"

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/130.0 Safari/537.36")

# Fráze v popisu, které stojí za pozornost (hledá se v textu bez diakritiky).
VAROVANI = {
    "Zákaz kouření": r"nekur|nekour|zakaz\w* koureni|koureni\w* (je )?(zakaz|neni povol|neni mozn)"
                     r"|nelze kourit|bez koureni|non-?smok",
    "Kouření OK": r"kuraci? (vitan|povolen|ok)|koureni (je )?(povoleno|mozne)",
    "Asi max. 2 osoby": r"(max\.?|maximalne|nejvyse|idealne|pouze|jen) ?(pro )?(1\s*[-–]\s*2|2|dve|dva|jednu|1) "
                        r"?(osob|lid|najemnik)|pro (jednu|1|1\s*[-–]\s*2) osob|\bpro (par|pary)\b"
                        r"|(vhodn\w*|ideal\w*) (pro|i pro) (par|jednotliv)",
    "Bez zvířat": r"bez (domacich )?(zvirat|mazlick)|zvirata (ne|nejsou)|zakaz zvirat",
}


def log(msg):
    line = f"{datetime.now():%Y-%m-%d %H:%M:%S}  {msg}"
    print(line, flush=True)
    with open(LOG_SOUBOR, "a", encoding="utf-8") as f:
        f.write(line + "\n")


def session():
    s = requests.Session()
    s.headers.update({"User-Agent": UA, "Accept-Language": "cs"})
    retry = Retry(total=4, backoff_factor=2, status_forcelist=(429, 500, 502, 503, 504),
                  allowed_methods=None)
    s.mount("https://", HTTPAdapter(max_retries=retry))
    return s


def bez_diakritiky(text):
    t = unicodedata.normalize("NFKD", text or "").lower()
    return "".join(c for c in t if not unicodedata.combining(c))


def vzdalenost_km(lat, lon):
    if lat is None or lon is None:
        return None
    la1, lo1, la2, lo2 = map(math.radians, (PRACE[0], PRACE[1], lat, lon))
    a = math.sin((la2 - la1) / 2) ** 2 + math.cos(la1) * math.cos(la2) * math.sin((lo2 - lo1) / 2) ** 2
    return round(6371 * 2 * math.asin(math.sqrt(a)), 1)


def v_dosahu(km):
    return MAX_VZDALENOST_KM is None or km is None or km <= MAX_VZDALENOST_KM


def najdi_varovani(text):
    t = re.sub(r"\s+", " ", bez_diakritiky(text))
    nalezy = []
    for nazev, vzor in VAROVANI.items():
        m = re.search(vzor, t)
        if m:
            a, b = max(0, m.start() - 30), min(len(t), m.end() + 30)
            nalezy.append(f"{nazev}: „…{t[a:b].strip()}…“")
    return "\n".join(nalezy)


CASTKA = re.compile(r"(?<![\d.,/])(\d{1,2}[ .]\d{3}|\d{3,5})(?![\d]|[.,]\d)")
KLIC_SLUZBY = re.compile(r"poplat|sluzb|zaloh|naklady na bydleni|energi|vyuctov")
KLIC_ELEKTRINA = re.compile(r"elektr|\bel\.")
NE_POPLATKY = re.compile(r"kauc|proviz|jistot|dph|garaz|parkov|stani|najemn[eiy]|najmu|najem\b|cena")
JEDNOTKA_NE = re.compile(r"\s*(m\b|m2|m²|km|min|metr|let\b|rok|%|\.? ?patr|mb|gb|kw|ks)")
ZA_OSOBU = re.compile(r"^\W*(kc|czk)?\W*(mes\w*\W*)?(poplat\w* |sluzb\w* )?(/|na |za |kazd\w* )\s*os")
PRO_N_OSOB = re.compile(r"(\b[12] ?os|jedn\w* osob|dve osob|dva lid)")


def _najdi_castku(t, klic, r, vynech=None):
    """První částka v rozumném rozsahu, která stojí u klíčového slova (za ním, nebo těsně před ním)."""
    for m in CASTKA.finditer(t):
        if vynech is not None and m.start() == vynech:
            continue
        pred, za = t[max(0, m.start() - 40):m.start()], t[m.end():m.end() + 40]
        k = None
        for k in klic.finditer(pred):
            pass                                      # poslední klíčové slovo před částkou
        if k and not NE_POPLATKY.search(pred[k.end():]):
            pass
        elif not klic.search(za[:18]) or NE_POPLATKY.search(pred[-25:]):  # "4 500 Kč služby"
            continue
        kc = int(re.sub(r"\D", "", m.group(1)))
        ma_kc = re.match(r"\s*(,-+)?\s*(kc|czk|,-)", za)
        if (not 300 <= kc <= 15_000 or kc == r["cena"] or JEDNOTKA_NE.match(za)
                or (1900 <= kc <= 2100 and not ma_kc)):
            continue
        return m, kc, pred, za
    return None


def odhad_poplatku(r):
    """Vrátí (poplatky Kč za měsíc pro POCET_OSOB, popis zdroje). Poplatky None = neznámé."""
    if r["poplatky"]:
        return r["poplatky"], "z inzerátu"
    for text in (r["pozn_cena"], r["popis"]):
        t = re.sub(r"\s+", " ", bez_diakritiky(text))
        nalez = _najdi_castku(t, KLIC_SLUZBY, r)
        if not nalez:
            continue
        m, kc, pred, za = nalez
        info = "vyčteno z textu"
        if ZA_OSOBU.match(za):
            kc, info = kc * POCET_OSOB, f"vyčteno z textu, {kc} Kč/osoba × {POCET_OSOB}"
        else:
            n = PRO_N_OSOB.search(pred[-30:] + za[:30])
            if n:
                info = f"vyčteno z textu – uvedeno „{n.group(1)}“, pro {POCET_OSOB} bude víc"
        # Elektřina uvedená zvlášť (ne "včetně elektřiny") se přičte.
        el = _najdi_castku(t, KLIC_ELEKTRINA, r, vynech=m.start())
        if el and not re.search(r"vcet\w* (el|energ)", el[2][-30:]) and el[1] <= 5_000:
            kc += el[1]
            info += f" + elektřina {el[1]}"
        return kc, info
    return None, "poplatky neuvedeny – ověřit"


def ts_datum(unix):
    return datetime.fromtimestamp(unix).strftime("%Y-%m-%d") if unix else None


# ---------------------------------------------------------------- SREALITY --
def sreality(s, zname):
    api = "https://www.sreality.cz/api/v1/estates"
    vysledky, offset, celkem = [], 0, None
    while celkem is None or offset < celkem:
        r = s.get(f"{api}/search", timeout=30, params=dict(
            category_main_cb=1, category_type_cb=2, category_sub_cb=4, locality_region_id=10,
            price_to=MAX_CENA, limit=100, offset=offset, sort="-date"))
        r.raise_for_status()
        d = r.json()
        celkem = d["pagination"]["total"]
        if not d["results"]:
            break
        vysledky += d["results"]
        offset += len(d["results"])

    out = []
    for e in vysledky:
        loc = e.get("locality") or {}
        km = vzdalenost_km(loc.get("gps_lat"), loc.get("gps_lon"))
        if not v_dosahu(km):
            continue
        hid = str(e["hash_id"])
        b = dict(
            id=f"sreality:{hid}", zdroj="Sreality",
            url=f"https://www.sreality.cz/detail/pronajem/byt/2+kk/praha/{hid}",
            nazev=e.get("advert_name"), cena=e.get("price_czk"),
            adresa=", ".join(x for x in (loc.get("street"), loc.get("citypart"), loc.get("district")) if x),
            lat=loc.get("gps_lat"), lon=loc.get("gps_lon"), km=km,
            metro_m=e.get("poi_metro_distance"),
        )
        if hid not in zname:          # detail jen u nových inzerátů
            time.sleep(PAUZA_S)
            try:
                dt = s.get(f"{api}/{hid}", timeout=30).json()["result"]
                b.update(
                    plocha=dt.get("usable_area") or dt.get("floor_area"),
                    patro=dt.get("floor_number"),
                    vybaveni=(dt.get("furnished") or {}).get("name"),
                    kauce=dt.get("refundable_deposit"),
                    od=dt.get("ready_date"),
                    pozn_cena=" ".join(str(x) for x in (dt.get("price_note"), dt.get("cost_of_living")) if x),
                    vlozeno=dt.get("since"),
                    varovani=najdi_varovani(dt.get("advert_description")),
                    popis=dt.get("advert_description"),
                )
            except Exception as ex:  # noqa: BLE001
                log(f"  Sreality detail {hid} selhal: {ex}")
        out.append(b)
    return out


# ------------------------------------------------------------- BEZREALITKY --
BZ_QUERY = """query($offset:Int){ listAdverts(offerType:[PRONAJEM], estateType:[BYT],
  disposition:[DISP_2_KK], priceTo:%d, regionOsmIds:["R435514"], currency:CZK,
  limit:60, offset:$offset, order:TIMEORDER_DESC){
  totalCount list{ id uri surface price charges deposit etage equipped availableFrom
  gps{lat lng} address(locale:CS) roommate reserved description } } }""" % MAX_CENA

VYBAVENI_BZ = {"VYBAVENY": "Ano", "CASTECNE": "Částečně", "NEVYBAVENY": "Ne"}


def bezrealitky(s, zname):
    out, offset, celkem = [], 0, None
    hdr = {"Origin": "https://www.bezrealitky.cz", "Referer": "https://www.bezrealitky.cz/"}
    while celkem is None or offset < celkem:
        r = s.post("https://api.bezrealitky.cz/graphql/", headers=hdr, timeout=30,
                   json=dict(query=BZ_QUERY, variables=dict(offset=offset)))
        r.raise_for_status()
        d = r.json()["data"]["listAdverts"]
        celkem = d["totalCount"]
        if not d["list"]:
            break
        offset += len(d["list"])
        for a in d["list"]:
            if a.get("roommate"):
                continue
            gps = a.get("gps") or {}
            km = vzdalenost_km(gps.get("lat"), gps.get("lng"))
            if not v_dosahu(km):
                continue
            out.append(dict(
                id=f"bezrealitky:{a['id']}", zdroj="Bezrealitky",
                url=f"https://www.bezrealitky.cz/nemovitosti-byty-domy/{a['uri']}",
                nazev=f"Pronájem bytu 2+kk {a.get('surface') or ''} m²" + (" (REZERVOVÁNO)" if a.get("reserved") else ""),
                cena=a.get("price"), poplatky=a.get("charges"), plocha=a.get("surface"),
                adresa=a.get("address"), lat=gps.get("lat"), lon=gps.get("lng"), km=km,
                patro=a.get("etage"), vybaveni=VYBAVENI_BZ.get(a.get("equipped"), a.get("equipped")),
                kauce=a.get("deposit"), od=ts_datum(a.get("availableFrom")),
                varovani=najdi_varovani(a.get("description")), popis=a.get("description"),
            ))
    return out


# ------------------------------------------------------------------ BAZOŠ --
def bazos(s, zname):
    base = "https://reality.bazos.cz"
    params = dict(hledat="2+kk", hlokalita="17100", cenado=MAX_CENA,
                  humkreis=(MAX_VZDALENOST_KM or 25) + 3)
    out, offset = [], 0
    while True:
        cesta = "/pronajmu/byt/" + (f"{offset}/" if offset else "")
        r = s.get(base + cesta, params=params, timeout=30)
        if r.status_code == 404 and offset:   # za poslední stránkou vrací Bazoš 404
            break
        r.raise_for_status()
        inz = BeautifulSoup(r.text, "lxml").select("div.inzeraty")
        if not inz:
            break
        for i in inz:
            a = i.select_one("h2.nadpis a")
            nazev = a.get_text(strip=True)
            if not re.search(r"2\s*\+?\s*kk", nazev, re.I):
                continue
            m = re.search(r"/inzerat/(\d+)/", a["href"])
            if not m:
                continue
            bid = m.group(1)
            cena_txt = i.select_one("div.inzeratycena").get_text(" ", strip=True)
            cena = int(re.sub(r"\D", "", cena_txt)) if re.search(r"\d", cena_txt) else None
            lok = i.select_one("div.inzeratylok").get_text(" ", strip=True)
            datum = re.search(r"\[(\d+)\.(\d+)\.\s*(\d{4})\]", i.get_text())
            plocha = re.search(r"(\d{2,3})\s*m(2|²)", nazev)
            b = dict(
                id=f"bazos:{bid}", zdroj="Bazoš", url=base + a["href"], nazev=nazev, cena=cena,
                adresa=lok, plocha=int(plocha.group(1)) if plocha else None,
                vlozeno=f"{datum.group(3)}-{int(datum.group(2)):02d}-{int(datum.group(1)):02d}" if datum else None,
            )
            if bid not in zname:
                time.sleep(PAUZA_S)
                try:
                    dt = s.get(b["url"], timeout=30).text
                    g = re.search(r"maps/place/([\d.]+),([\d.]+)", dt)
                    if g:
                        b["lat"], b["lon"] = float(g.group(1)), float(g.group(2))
                        b["km"] = vzdalenost_km(b["lat"], b["lon"])
                    p = BeautifulSoup(dt, "lxml").select_one("div.popisdetail")
                    if p:
                        b["popis"] = p.get_text("\n", strip=True)
                        b["varovani"] = najdi_varovani(b["popis"])
                except Exception as ex:  # noqa: BLE001
                    log(f"  Bazoš detail {bid} selhal: {ex}")
            out.append(b)
        offset += 20
        time.sleep(PAUZA_S)
    return out


# --------------------------------------------------------------- DATABÁZE --
SLOUPCE = ["zdroj", "url", "nazev", "cena", "poplatky", "plocha", "adresa", "lat", "lon", "km",
           "metro_m", "patro", "vybaveni", "kauce", "od", "pozn_cena", "vlozeno", "varovani", "popis"]


def otevri_db():
    db = sqlite3.connect(DB_SOUBOR)
    db.row_factory = sqlite3.Row
    db.execute(f"""CREATE TABLE IF NOT EXISTS byty (id TEXT PRIMARY KEY,
        {", ".join(c + " " for c in SLOUPCE)}, cena_puvodni, prvni_videni TEXT, posledni_videni TEXT,
        aktivni INTEGER DEFAULT 1)""")
    return db


def uloz(db, byty, beh):
    nove = 0
    for b in byty:
        stary = db.execute("SELECT * FROM byty WHERE id=?", (b["id"],)).fetchone()
        if stary is None:
            nove += 1
            cols = ["id"] + SLOUPCE + ["prvni_videni", "posledni_videni", "aktivni"]
            vals = [b["id"]] + [b.get(c) for c in SLOUPCE] + [beh, beh, 1]
            db.execute(f"INSERT INTO byty ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})", vals)
        else:
            # Aktualizuj jen to, co přišlo (detail se u starých nestahuje znovu).
            zmeny = {c: b[c] for c in SLOUPCE if b.get(c) is not None}
            if stary["cena"] and b.get("cena") and b["cena"] != stary["cena"]:
                zmeny["cena_puvodni"] = stary["cena_puvodni"] or stary["cena"]
            zmeny.update(posledni_videni=beh, aktivni=1)
            db.execute(f"UPDATE byty SET {', '.join(k + '=?' for k in zmeny)} WHERE id=?",
                       [*zmeny.values(), b["id"]])
    return nove


# ------------------------------------------------------------------ EXCEL --
HLAVICKA = [("Zdroj", "zdroj", 12), ("Inzerát", "nazev", 38), ("Celkem", "celkem", 10),
            ("Nájem", "cena", 10), ("Poplatky", "popl", 10), ("Poplatky – odkud", "popl_info", 30),
            ("m²", "plocha", 6), ("km od Zoo", "km", 10), ("Adresa / lokalita", "adresa", 34),
            ("Metro (m)", "metro_m", 10), ("Patro", "patro", 6), ("Vybavení", "vybaveni", 10),
            ("Kauce", "kauce", 10), ("Volné od", "od", 11), ("Pozor v popisu", "varovani", 50),
            ("Pozn. k ceně", "pozn_cena", 28), ("Dřív stál", "cena_puvodni", 10), ("Vloženo", "vlozeno", 11),
            ("Poprvé viděn", "prvni_videni", 16)]
SLOUPEC = {klic: i for i, (_, klic, _) in enumerate(HLAVICKA, 1)}

CERVENA = PatternFill("solid", fgColor="F8D7DA")
ORANZOVA = PatternFill("solid", fgColor="FFE8C2")
ZELENA = PatternFill("solid", fgColor="D4EDDA")
ZLUTA = PatternFill("solid", fgColor="FFF3B0")


def s_cenou(radky):
    """Dopočítá poplatky a celkovou cenu, vyřadí byty dražší než MAX_CENA."""
    out = []
    for r in map(dict, radky):
        r["popl"], r["popl_info"] = odhad_poplatku(r)
        r["celkem"] = r["cena"] + r["popl"] if r["cena"] and r["popl"] is not None else None
        if (r["cena"] or 0) > MAX_CENA or (r["celkem"] or 0) > MAX_CENA:
            continue
        out.append(r)
    out.sort(key=lambda r: (r["km"] is None, r["km"] or 0, r["celkem"] or r["cena"] or 0))
    return out


def list_excelu(ws, radky):
    ws.append([h for h, _, _ in HLAVICKA])
    for c, (_, _, w) in enumerate(HLAVICKA, 1):
        ws.cell(1, c).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, c).fill = PatternFill("solid", fgColor="34495E")
        ws.column_dimensions[ws.cell(1, c).column_letter].width = w
    for r in radky:
        ws.append([r[klic] for _, klic, _ in HLAVICKA])
        n = ws.max_row
        bunka = lambda klic: ws.cell(n, SLOUPEC[klic])  # noqa: E731
        bunka("nazev").hyperlink = r["url"]
        bunka("nazev").font = Font(color="0563C1", underline="single")
        for k in ("celkem", "cena", "popl", "kauce", "cena_puvodni"):
            bunka(k).number_format = "#,##0"
        bunka("celkem").font = Font(bold=True)
        if r["popl"] is None or "1 osobu" in r["popl_info"]:
            bunka("popl_info").fill = ZLUTA
        var = r["varovani"] or ""
        if "Zákaz kouření" in var:
            bunka("varovani").fill = CERVENA
        elif "max. 2" in var:
            bunka("varovani").fill = ORANZOVA
        elif "Kouření OK" in var:
            bunka("varovani").fill = ZELENA
        bunka("varovani").alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def uloz_excel(db, beh):
    # Bazoš hledá v kruhu podle PSČ, přesná vzdálenost je známa až z detailu -> dofiltrovat tady.
    dosah = "" if MAX_VZDALENOST_KM is None else f"AND (km IS NULL OR km <= {MAX_VZDALENOST_KM})"
    nove = s_cenou(db.execute(f"SELECT * FROM byty WHERE prvni_videni=? {dosah}", (beh,)))
    vse = s_cenou(db.execute(f"SELECT * FROM byty WHERE aktivni=1 {dosah}"))
    wb = Workbook()
    ws = wb.active
    ws.title = f"Nové ({len(nove)})"
    list_excelu(ws, nove)
    list_excelu(wb.create_sheet(f"Všechny ({len(vse)})"), vse)
    cil = EXCEL_SOUBOR
    try:
        wb.save(cil)
    except PermissionError:   # Excel má soubor otevřený
        cil = SLOZKA / f"byty_{datetime.now():%Y%m%d_%H%M}.xlsx"
        wb.save(cil)
        log(f"byty.xlsx je otevřený v Excelu, uloženo jako {cil.name}")
    return cil, len(nove), len(vse)


# -------------------------------------------------------------------- WEB --
WEB_POLE = ["id", "zdroj", "url", "nazev", "celkem", "cena", "popl", "popl_info", "plocha", "km",
            "adresa", "lat", "lon", "metro_m", "patro", "vybaveni", "kauce", "od", "varovani",
            "pozn_cena", "cena_puvodni", "vlozeno", "prvni_videni", "popis"]


def cislo(x):
    return int(x) if isinstance(x, float) and x.is_integer() else x


def uloz_web(db, beh):
    dosah = "" if MAX_VZDALENOST_KM is None else f"AND (km IS NULL OR km <= {MAX_VZDALENOST_KM})"
    byty = []
    for r in s_cenou(db.execute(f"SELECT * FROM byty WHERE aktivni=1 {dosah}")):
        b = {k: cislo(r.get(k)) for k in WEB_POLE}
        if b["popis"] and len(b["popis"]) > 2500:
            b["popis"] = b["popis"][:2500] + "…"
        byty.append({k: v for k, v in b.items() if v not in (None, "")})
    data = dict(aktualizovano=beh, max_cena=MAX_CENA, osob=POCET_OSOB, max_km=MAX_VZDALENOST_KM,
                prace=PRACE, byty=byty)
    WEB_JSON.parent.mkdir(exist_ok=True)
    WEB_JSON.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    return len(byty)


# ------------------------------------------------------------------- MAIN --
def main():
    beh = datetime.now().strftime("%Y-%m-%d %H:%M")
    log(f"=== Start ({beh}) ===")
    s, db = session(), otevri_db()
    for nazev, fce in (("Sreality", sreality), ("Bezrealitky", bezrealitky), ("Bazoš", bazos)):
        try:
            zname = {r["id"].split(":", 1)[1] for r in db.execute(
                "SELECT id FROM byty WHERE zdroj=? AND popis IS NOT NULL", (nazev,))}
            byty = fce(s, zname)
            nove = uloz(db, byty, beh)
            # Co zdroj už nevrací, je pryč (pronajato/smazáno). Jen když zdroj prošel bez chyby.
            db.execute("UPDATE byty SET aktivni=0 WHERE zdroj=? AND posledni_videni<>?", (nazev, beh))
            db.commit()
            log(f"{nazev}: {len(byty)} v dosahu, z toho {nove} nových")
        except Exception as ex:  # noqa: BLE001
            db.rollback()
            log(f"{nazev}: CHYBA {type(ex).__name__}: {ex}")
    log(f"Web: {WEB_JSON}  ({uloz_web(db, beh)} bytů)")
    if not os.environ.get("CI"):
        cil, n_nove, n_vse = uloz_excel(db, beh)
        log(f"Excel: {cil}  (nové: {n_nove}, aktivní celkem: {n_vse})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
